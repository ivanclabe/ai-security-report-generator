"""Builds the .docx from a validated ReportSpecification, reproducing the layout of
the client's reference "Informe de Gestión CSOC" (form GCB-FR-14).

Purely presentational: it draws what the specification contains -- every value
was computed upstream from the data (csoc_facts.py / templates.py). Layout of the
reference that is reproduced:

* page 1: cover -- title, client logo (or name), period, sweeping green/teal art
* page 2: "Tabla de Contenido" with the vertical title and the chapter index
* page 3: document control table (title, author, version, review/approval)
* page 4: 1. OBJETIVO (with the ringed graphic) and 2. SERVICIOS (service tiles)
* from chapter 3: gray header band with the Invotecsa logo + "Informe de Gestión SOC"
* every page but the cover: the footer box CÓDIGO / PÁGINA x de y / VERSIÓN / FECHA / SECRETO
* green table headers, light-green value cells, colored risk cells, gray totals,
  boxed significant alarms, and charts in the same colors.

Deliberately not reproduced: stock photography and the reference's cyan/yellow
highlighter (removed earlier at the client's request -- "texto limpio y sin resaltados").
"""
import io
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.document import report_art
from app.document.font_embedding import embed_font
from app.schemas.report_spec import (
    AlertBoxBlock,
    BulletsBlock,
    ChartBlock,
    MinorHeadingBlock,
    NoticeBlock,
    ParagraphBlock,
    ReportSpecification,
    Section,
    ServiceTilesBlock,
    SideBySideBlock,
    SimpleTable,
    Subsection,
    TableBlock,
)

ASSETS_DIR = Path(__file__).resolve().parent / "assets"
LOGOS_DIR = ASSETS_DIR / "logos"
DEFAULT_LOGO_PATH = ASSETS_DIR / "invotecsa_logo.png"
_LOGO_EXTENSIONS = (".png", ".jpg", ".jpeg")

HEADING_FONT = "Poppins"  # embedded in the document (font_embedding.py)
BODY_FONT = "Calibri"  # Office's default; servers render it with the metric-compatible Carlito
DARK_GREEN = RGBColor(0x00, 0x53, 0x2A)
HEADER_FILL = "0B5D33"
MINOR_GREEN = RGBColor(0x0B, 0x5D, 0x33)
TEXT = RGBColor(0x3A, 0x3A, 0x3A)
MUTED = RGBColor(0x5B, 0x66, 0x77)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEAL_HEX = "1A9C96"
LABEL_FILL = "D9D9D9"
RISK_FILL = {"Critico": "C00000", "Alto": "FF0000", "Medio": "FFC000", "Bajo": "00B050"}
RISK_LABEL = {"Critico": "Crítico", "Alto": "Alto", "Medio": "Medio", "Bajo": "Bajo"}

PAGE_WIDTH = 8.5
MARGIN_SIDE = 1.0
CONTENT_WIDTH = PAGE_WIDTH - 2 * MARGIN_SIDE
NO_EVIDENCE = "No se encontró evidencia suficiente en las fuentes de datos disponibles para completar esta sección."


def resolve_logo_path(logo_key: str | None) -> Path | None:
    """Logo file for a company: assets/logos/<slug of the logo key>.png|jpg|jpeg
    (e.g. "Thomas TPS" -> logos/thomas-tps.png), falling back to the default
    logo when that company's file has not been supplied yet. The slug keeps only
    [a-z0-9-], so a logo key can never point outside the logos directory."""
    if logo_key:
        slug = re.sub(r"[^a-z0-9]+", "-", logo_key.lower()).strip("-")
        for extension in _LOGO_EXTENSIONS:
            candidate = LOGOS_DIR / f"{slug}{extension}"
            if slug and candidate.is_file():
                return candidate
    return DEFAULT_LOGO_PATH if DEFAULT_LOGO_PATH.exists() else None


def _company_logo(logo_key: str | None) -> Path | None:
    """The CLIENT's own logo for the cover, or None (never the provider's logo in its place)."""
    path = resolve_logo_path(logo_key)
    return path if path and path != DEFAULT_LOGO_PATH else None


# ---- low-level helpers -------------------------------------------------------------------------


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.upper())


# Word rejects documents whose property elements are out of the schema's order, so every
# element is inserted at its place in these sequences (ECMA-376 CT_TcPr / CT_TrPr / CT_TblPr).
_TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
               "textDirection", "tcFitText", "vAlign", "hideMark"]
_TRPR_ORDER = ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit", "trHeight",
               "tblHeader", "tblCellSpacing", "jc", "hidden"]
_TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW",
                "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]


def _put(parent, element, order: list[str]) -> None:
    """Replace parent's child of the same tag, keeping the schema order."""
    tag = element.tag.split("}")[1]
    for old in parent.findall(qn(f"w:{tag}")):
        parent.remove(old)
    position = order.index(tag)
    for index, child in enumerate(parent):
        child_tag = child.tag.split("}")[1] if "}" in str(child.tag) else ""
        if child_tag in order and order.index(child_tag) > position:
            parent.insert(index, element)
            return
    parent.append(element)


def _shade(cell, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    _put(cell._tc.get_or_add_tcPr(), shd, _TCPR_ORDER)


def _widths(table, widths: list[float]) -> None:
    """Fixed column widths that Word and LibreOffice both honor (grid + cell widths + table width)."""
    table.autofit = False
    tbl = table._tbl
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    _put(tbl.tblPr, tbl_w, _TBLPR_ORDER)
    for grid_col, width in zip(tbl.tblGrid.findall(qn("w:gridCol")), widths, strict=False):
        grid_col.set(qn("w:w"), str(int(width * 1440)))
    for tr in tbl.tr_lst:
        col = 0
        for tc in tr.tc_lst:
            span = tc.grid_span
            tc_w = OxmlElement("w:tcW")
            tc_w.set(qn("w:w"), str(int(sum(widths[col:col + span]) * 1440)))
            tc_w.set(qn("w:type"), "dxa")
            _put(tc.get_or_add_tcPr(), tc_w, _TCPR_ORDER)
            col += span


def _font(run, *, name: str | None = None, size: float | None = None, bold: bool | None = None,
          color: RGBColor | None = None, italic: bool | None = None) -> None:
    if name:
        run.font.name = name
        r_pr = run._r.get_or_add_rPr()
        fonts = r_pr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            r_pr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _cell_text(cell, text: str, *, bold=False, color: RGBColor | None = None, size: float = 10.5,
               align: str = "left", font: str = BODY_FONT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                           "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    _font(paragraph.add_run(text), name=font, size=size, bold=bold, color=color)


def _no_borders(table) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    _put(table._tbl.tblPr, borders, _TBLPR_ORDER)


def _cell_border(cell, edge: str, size_eighths: int, color: str) -> None:
    borders = OxmlElement("w:tcBorders")
    element = OxmlElement(f"w:{edge}")
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), str(size_eighths))
    element.set(qn("w:color"), color)
    borders.append(element)
    _put(cell._tc.get_or_add_tcPr(), borders, _TCPR_ORDER)


def _repeat_header(row) -> None:
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    _put(row._tr.get_or_add_trPr(), element, _TRPR_ORDER)


def _cant_split(row) -> None:
    _put(row._tr.get_or_add_trPr(), OxmlElement("w:cantSplit"), _TRPR_ORDER)


def _add_field(paragraph, instruction: str, *, size: float | None = None, bold: bool = False, placeholder: str = "") -> None:
    def run_with(element):
        run = paragraph.add_run()
        _font(run, size=size, bold=bold)
        run._r.append(element)
        return run

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run_with(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    run_with(instr)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run_with(separate)
    if placeholder:
        _font(paragraph.add_run(placeholder), size=size, bold=bold)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run_with(end)


_bookmark_id = [0]


def _bookmark(paragraph, name: str) -> None:
    _bookmark_id[0] += 1
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(_bookmark_id[0]))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(_bookmark_id[0]))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _bookmark_name(number: str) -> str:
    return "_Toc_s" + number.replace(".", "_")


def _picture(container, png: bytes, width_in: float, *, align=WD_ALIGN_PARAGRAPH.CENTER, bleed: bool = False):
    paragraph = container.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(4)
    if bleed:
        paragraph.paragraph_format.left_indent = Inches(-MARGIN_SIDE)
        paragraph.paragraph_format.right_indent = Inches(-MARGIN_SIDE)
    paragraph.add_run().add_picture(io.BytesIO(png), width=Inches(width_in))
    return paragraph


# ---- styles / page ------------------------------------------------------------------------------


def _style_document(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2

    for name, size in (("Heading 1", 20), ("Heading 2", 14)):
        style = styles[name]
        style.font.name = HEADING_FONT
        style.element.rPr.rFonts.set(qn("w:asciiTheme"), "")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = DARK_GREEN
        style.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    section = document.sections[0]
    section.page_width = Inches(PAGE_WIDTH)
    section.page_height = Inches(11)
    section.left_margin = section.right_margin = Inches(MARGIN_SIDE)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.55)
    section.footer_distance = Inches(0.35)


def _footer_box(footer, spec: ReportSpecification) -> None:
    """The client template's control box: CÓDIGO/PÁGINA, value row (page x de y),
    VERSIÓN/FECHA, value row, classification across both columns."""
    dc = spec.document_control
    for paragraph in footer.paragraphs:
        paragraph.text = ""
    table = footer.add_table(rows=5, cols=2, width=Inches(2.55))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = [("CÓDIGO", "PÁGINA", True), (dc.codigo_documento, None, False), ("VERSIÓN", "FECHA", True),
            (dc.version_plantilla, dc.fecha_plantilla, False)]
    for i, (left, right, bold) in enumerate(rows):
        cells = table.rows[i].cells
        _cell_text(cells[0], left, bold=bold, size=7.5 if bold else 7, align="center", font="Calibri")
        if right is None:
            cells[1].text = ""
            paragraph = cells[1].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            _add_field(paragraph, "PAGE", size=7, placeholder="1")
            _font(paragraph.add_run(" de "), size=7)
            _add_field(paragraph, "NUMPAGES", size=7, placeholder="1")
        else:
            _cell_text(cells[1], right, bold=bold, size=7.5 if bold else 7, align="center", font="Calibri")
    merged = table.rows[4].cells[0].merge(table.rows[4].cells[1])
    _cell_text(merged, dc.clasificacion, bold=True, size=7.5, align="center", font="Calibri")
    _widths(table, [1.27, 1.27])


def _header_band(section) -> None:
    section.header_distance = Inches(0)
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.text = ""
    paragraph.paragraph_format.left_indent = Inches(-MARGIN_SIDE)
    paragraph.paragraph_format.right_indent = Inches(-MARGIN_SIDE)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_picture(io.BytesIO(report_art.header_band()), width=Inches(PAGE_WIDTH))


# ---- front matter ----------------------------------------------------------------------------------


def _cover(document: Document, spec: ReportSpecification) -> None:
    for _ in range(3):
        document.add_paragraph()
    for line in ("Informe de Gestión", "Servicio de Monitoreo SOC"):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        _font(paragraph.add_run(line), name=HEADING_FONT, size=27, bold=True, color=DARK_GREEN)
    logo = _company_logo(spec.company_logo)
    if logo:
        _picture(document, logo.read_bytes(), 2.6)
    else:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(14)
        _font(paragraph.add_run(spec.company), name=HEADING_FONT, size=24, color=RGBColor(0x4D, 0x4D, 0x4D))
    period = document.add_paragraph()
    period.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(period.add_run(_period_text(spec)), name="Calibri", size=13, color=TEXT)
    art = _picture(document, report_art.cover_art(PAGE_WIDTH, 4.7), PAGE_WIDTH, bleed=True)
    art.paragraph_format.space_before = Pt(10)
    art.add_run().add_break(WD_BREAK.PAGE)


def _period_text(spec: ReportSpecification) -> str:
    from datetime import date

    from app.services.report.csoc_facts import spanish_date

    start, end = date.fromisoformat(spec.period.start), date.fromisoformat(spec.period.end)
    return f"{spanish_date(start)} al {spanish_date(end)}"


def _toc_entries(spec: ReportSpecification) -> list[tuple[str, str, int]]:
    entries = [("1", "OBJETIVO", 1), ("2", "SERVICIOS", 1)]
    for section in spec.sections:
        entries.append((section.number, section.title.upper(), 1))
        entries += [(sub.number, sub.title, 2) for sub in section.subsections]
    entries.append(("11", "RECOMENDACIONES", 1))
    return entries


def _table_of_contents(document: Document, spec: ReportSpecification) -> None:
    table = document.add_table(rows=1, cols=2)
    _no_borders(table)
    left, right = table.rows[0].cells
    _widths(table, [1.1, CONTENT_WIDTH - 1.1])
    table.rows[0].height = Inches(7.6)
    _cell_border(left, "right", 36, TEAL_HEX)
    direction = OxmlElement("w:textDirection")
    direction.set(qn("w:val"), "btLr")
    _put(left._tc.get_or_add_tcPr(), direction, _TCPR_ORDER)
    paragraph = left.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(paragraph.add_run("Tabla de Contenido"), name=HEADING_FONT, size=44, bold=True, color=DARK_GREEN)

    first = True
    for number, title, level in _toc_entries(spec):
        entry = right.paragraphs[0] if first else right.add_paragraph()
        first = False
        fmt = entry.paragraph_format
        fmt.space_after = Pt(5 if level == 1 else 3)
        fmt.space_before = Pt(6 if level == 1 else 0)
        fmt.left_indent = Inches(0.25 if level == 1 else 0.65)
        fmt.tab_stops.add_tab_stop(Inches(CONTENT_WIDTH - 1.25), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        label = f"{number}.\t{title}" if level == 1 else f"{number} {title}"
        if level == 1:
            fmt.tab_stops.add_tab_stop(Inches(0.75))
        _font(entry.add_run(label), name=HEADING_FONT, size=10.5 if level == 1 else 9.5, bold=True, color=TEXT)
        _font(entry.add_run("\t"), size=10)
        _add_field(entry, f"PAGEREF {_bookmark_name(number)} \\h", size=10, placeholder="1")
    # no break here: the document-control page starts with "page break before", which never
    # produces a blank page when the index table already fills this one


def _document_control(document: Document, spec: ReportSpecification) -> None:
    dc = spec.document_control
    for i in range(14):
        document.add_paragraph().paragraph_format.page_break_before = i == 0
    table = document.add_table(rows=6, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    title_cell = table.rows[0].cells[1].merge(table.rows[0].cells[3])
    _cell_text(table.rows[0].cells[0], "Título", color=WHITE, align="center", font="Calibri", size=11)
    _cell_text(title_cell, dc.titulo, color=WHITE, align="center", font="Calibri", size=11)
    _shade(table.rows[0].cells[0], HEADER_FILL)
    _shade(title_cell, HEADER_FILL)
    author_cell = table.rows[1].cells[1].merge(table.rows[1].cells[3])
    _cell_text(table.rows[1].cells[0], "Autor", bold=True, align="center", font="Calibri", size=10.5)
    _cell_text(author_cell, dc.autor, align="center", size=10.5)
    rows = [
        ("Versión", f"{dc.version:02d}", "Fecha Versión", dc.fecha_version),
        ("¹Revisado por:", dc.revisado_por, "Fecha Revisión", dc.fecha_revision),
        ("²Revisado por:", dc.revisado_por_2, "Fecha Aprobación", dc.fecha_aprobacion),
        ("²Aprobado por:", dc.aprobado_por, "Nº Total Páginas", None),
    ]
    for i, (label, value, label2, value2) in enumerate(rows, start=2):
        cells = table.rows[i].cells
        _cell_text(cells[0], label, bold=True, align="center", font="Calibri", size=10.5)
        _cell_text(cells[1], value, align="center", size=10.5)
        _cell_text(cells[2], label2, bold=True, align="center", font="Calibri", size=10.5)
        if value2 is None:
            cells[3].text = ""
            paragraph = cells[3].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            _add_field(paragraph, "NUMPAGES", size=10.5, bold=True, placeholder="1")
        else:
            _cell_text(cells[3], value2, align="center", size=10.5)
    _widths(table, [1.3, 2.4, 1.35, 1.25])
    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    _font(note.add_run(dc.nota), size=9, italic=True, color=MUTED)
    document.add_page_break()


def _heading(container, text: str, level: int, number: str, *, align=None):
    paragraph = container.add_paragraph(style="Heading 1" if level == 1 else "Heading 2")
    if align is not None:
        paragraph.alignment = align
    _font(paragraph.add_run(text), name=HEADING_FONT)
    _bookmark(paragraph, _bookmark_name(number))
    return paragraph


def _objective_and_services(document: Document, spec: ReportSpecification) -> None:
    table = document.add_table(rows=1, cols=2)
    _no_borders(table)
    left, right = table.rows[0].cells
    _widths(table, [3.9, CONTENT_WIDTH - 3.9])
    left.paragraphs[0].paragraph_format.space_after = Pt(0)
    for _ in range(3):
        left.add_paragraph()
    _heading(left, "1. OBJETIVO", 1, "1")
    paragraph = left.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _font(paragraph.add_run(spec.objective), size=11)
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraphs[0].add_run().add_picture(io.BytesIO(report_art.objective_art()), width=Inches(2.45))

    document.add_paragraph()
    _heading(document, "2. SERVICIOS", 1, "2", align=WD_ALIGN_PARAGRAPH.RIGHT)
    if spec.services:
        from datetime import date

        from app.services.report.csoc_facts import period_label

        label = period_label(date.fromisoformat(spec.period.start), date.fromisoformat(spec.period.end))
        phrase = f"de {label}" if not label.startswith("del") else label
        _paragraph(document, f"A continuación, se listan los servicios administrados por el SOC INVOTECSA durante el "
                             f"período {phrase}, según las alarmas registradas en las fuentes de datos:")
        _picture(document, report_art.service_tiles(tuple(spec.services)), min(CONTENT_WIDTH, 1.6 * len(spec.services)))
    else:
        _notice(document, NO_EVIDENCE)


# ---- blocks -------------------------------------------------------------------------------------------


def _paragraph(container, text: str, bold_terms: list[str] | None = None) -> None:
    paragraph = container.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    terms = [t for t in (bold_terms or []) if t and t in text]
    if not terms:
        _font(paragraph.add_run(text), size=11)
        return
    pattern = "(" + "|".join(re.escape(t) for t in terms) + ")"
    for part in re.split(pattern, text):
        if part:
            _font(paragraph.add_run(part), size=11, bold=part in terms)


def _notice(container, text: str) -> None:
    paragraph = container.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _font(paragraph.add_run(text), size=10.5, italic=True, color=MUTED)


def _minor_heading(container, text: str) -> None:
    paragraph = container.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.keep_with_next = True
    _font(paragraph.add_run(text), name="Calibri", size=12.5, color=MINOR_GREEN)


def _bullets(container, items: list[str]) -> None:
    for item in items:
        paragraph = container.add_paragraph(style="List Bullet")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _font(paragraph.add_run(item), size=11)


def _table(container, block: TableBlock, *, width: float | None = None):
    """`width`: the space available (a nested table gets its cell's width)."""
    columns = sum(c.colspan for c in block.header_rows[0])
    table = container.add_table(rows=0, cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_count = len(block.header_rows)
    all_rows = [*block.header_rows, *block.rows]
    spans_above: dict[int, bool] = {}
    for r, row_spec in enumerate(all_rows):
        row = table.add_row()
        _cant_split(row)
        if r < header_count:
            _repeat_header(row)
        col = 0
        for cell_spec in row_spec:
            cell = row.cells[col]
            if cell_spec.colspan > 1:
                cell = cell.merge(row.cells[col + cell_spec.colspan - 1])
            is_header = r < header_count
            _cell_text(cell, cell_spec.text, bold=cell_spec.bold, size=block.font_size,
                       color=_rgb(cell_spec.color) if cell_spec.color else None, align=cell_spec.align,
                       font="Calibri" if is_header or cell_spec.bold else BODY_FONT)
            if cell_spec.fill:
                _shade(cell, cell_spec.fill)
            # an empty header cell under a single-column header is the same cell, spanning two rows
            if is_header and r > 0 and cell_spec.text == "" and spans_above.get(col):
                table.cell(r - 1, col).merge(cell)
            if is_header and r == 0:
                spans_above[col] = cell_spec.colspan == 1
            col += cell_spec.colspan
    _cell_margins(table, 35, 80)
    available = width or CONTENT_WIDTH
    widths = list(block.widths) if block.widths and len(block.widths) == columns else [available / columns] * columns
    if sum(widths) > available:
        widths = [w * available / sum(widths) for w in widths]
    _widths(table, widths)
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = 1  # center
    return table


def _chart(container, block: ChartBlock) -> None:
    png = report_art.chart(
        block.chart, block.title, block.categories,
        [{"name": s.name, "values": s.values, "color": s.color} for s in block.series],
        width=block.width_inches, height=block.height_inches, x_label=block.x_label, y_label=block.y_label,
    )
    _picture(container, png, min(block.width_inches, CONTENT_WIDTH))


def _cell_margins(table, vertical: int, horizontal: int) -> None:
    """Table-wide cell padding in twips (LibreOffice and Word otherwise draw text against the borders)."""
    margins = OxmlElement("w:tblCellMar")
    for edge, twips in (("top", vertical), ("bottom", vertical), ("left", horizontal), ("right", horizontal)):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), str(twips))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    _put(table._tbl.tblPr, margins, _TBLPR_ORDER)


def _alert_box(container, block: AlertBoxBlock) -> None:
    box = container.add_table(rows=4, cols=1)
    box.style = "Table Grid"
    _cell_margins(box, 60, 100)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    _widths(box, [CONTENT_WIDTH])
    title = box.rows[0].cells[0]
    _cell_text(title, block.alarm_number, bold=True, color=WHITE, align="center", font="Calibri", size=11)
    _shade(title, HEADER_FILL)
    _cell_text(box.rows[1].cells[0], block.description, align="justify", size=10.5)

    info_cell = box.rows[2].cells[0]
    info_cell.paragraphs[0].paragraph_format.space_after = Pt(2)
    info = info_cell.add_table(rows=1, cols=6)
    info.style = "Table Grid"
    _cell_margins(info, 20, 40)
    banner = info.rows[0].cells[0].merge(info.rows[0].cells[5])
    _cell_text(banner, "INFORMACIÓN DEL EVENTO DE SEGURIDAD O DISPONIBILIDAD", bold=True, color=WHITE, align="center",
               font="Calibri", size=8)
    _shade(banner, HEADER_FILL)
    pairs = list(block.event_info)
    for start in range(0, len(pairs), 3):
        cells = info.add_row().cells
        for j, (label, value) in enumerate(pairs[start:start + 3]):
            _cell_text(cells[j * 2], label, bold=True, align="center", font="Calibri", size=7.5)
            _shade(cells[j * 2], LABEL_FILL)
            risk = next((k for k, v in RISK_LABEL.items() if v == value), None) if label.lower().startswith("riesgo") else None
            _cell_text(cells[j * 2 + 1], value, align="center", size=7.5, bold=bool(risk),
                       color=WHITE if risk else None)
            if risk:
                _shade(cells[j * 2 + 1], RISK_FILL[risk])
    _widths(info, [0.95, 1.2, 0.95, 1.2, 0.95, 1.05])
    if block.recommendation:
        paragraph = info_cell.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _font(paragraph.add_run("Recomendación registrada: "), size=9.5, bold=True)
        _font(paragraph.add_run(block.recommendation), size=9.5)

    response = box.rows[3].cells[0]
    response.text = ""
    paragraph = response.paragraphs[0]
    _font(paragraph.add_run("Respuesta:"), name="Calibri", size=11, bold=True)
    body = response.add_paragraph()
    body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if block.response:
        _font(body.add_run(block.response), size=10.5)
        if block.responded_at:
            when = response.add_paragraph()
            _font(when.add_run(f"Fecha de respuesta: {block.responded_at}"), size=9.5, italic=True, color=MUTED)
    else:
        _font(body.add_run("Sin respuesta registrada en la fuente de datos al corte del informe."), size=10.5,
              italic=True, color=MUTED)
    container.add_paragraph()


def _side_by_side(container, block: SideBySideBlock) -> None:
    outer = container.add_table(rows=2, cols=2)
    outer.style = "Table Grid"
    outer.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, title in enumerate(block.titles):
        _cell_text(outer.rows[0].cells[i], title, color=WHITE, align="center", font="Calibri", size=11)
        _shade(outer.rows[0].cells[i], HEADER_FILL)
    _widths(outer, [CONTENT_WIDTH / 2, CONTENT_WIDTH / 2])
    for cell, table_block in ((outer.rows[1].cells[0], block.left), (outer.rows[1].cells[1], block.right)):
        cell.paragraphs[0].paragraph_format.space_after = Pt(2)
        _table(cell, table_block, width=CONTENT_WIDTH / 2 - 0.25)
    container.add_paragraph()


def _render_blocks(container, blocks) -> None:
    for block in blocks:
        if isinstance(block, ParagraphBlock):
            if block.style == "caption":
                paragraph = container.add_paragraph()
                _font(paragraph.add_run(block.text), size=9, italic=True, color=MUTED)
            else:
                _paragraph(container, block.text, block.bold_terms)
        elif isinstance(block, MinorHeadingBlock):
            _minor_heading(container, block.text)
        elif isinstance(block, BulletsBlock):
            _bullets(container, block.items)
        elif isinstance(block, NoticeBlock):
            _notice(container, block.text)
        elif isinstance(block, TableBlock):
            _table(container, block)
            container.add_paragraph()
        elif isinstance(block, ChartBlock):
            _chart(container, block)
        elif isinstance(block, AlertBoxBlock):
            _alert_box(container, block)
        elif isinstance(block, ServiceTilesBlock):
            _picture(container, report_art.service_tiles(tuple(block.services)), CONTENT_WIDTH)
        elif isinstance(block, SideBySideBlock):
            _side_by_side(container, block)


# ---- legacy section fields (specifications stored before the block layout) ------------------------


def _legacy_table(document, table_spec: SimpleTable) -> None:
    if not table_spec.rows:
        _notice(document, "Sin datos para el período.")
        return
    from app.schemas.report_spec import Cell

    block = TableBlock(
        header_rows=[[Cell(text=h, fill=HEADER_FILL, bold=True, color="FFFFFF", align="center") for h in table_spec.headers]],
        rows=[[Cell(text=str(v)) for v in row] for row in table_spec.rows],
    )
    _table(document, block)
    document.add_paragraph()


def _render_legacy(document, item) -> None:
    if not item.available:
        _notice(document, item.unavailable_reason or NO_EVIDENCE)
        return
    if item.summary:
        _paragraph(document, item.summary)
    _bullets(document, item.body)
    if item.table:
        _legacy_table(document, item.table)
    for alert in item.significant_alerts:
        _alert_box(document, AlertBoxBlock(alarm_number=alert.heading, description=alert.description,
                                           response=alert.response))


def _render_section(document: Document, section: Section) -> None:
    _heading(document, f"{section.number}. {section.title.upper()}", 1, section.number)
    _render_legacy(document, section)
    _render_blocks(document, section.blocks)
    for sub in section.subsections:
        _render_subsection(document, sub)


def _render_subsection(document: Document, sub: Subsection) -> None:
    _heading(document, f"{sub.number} {sub.title}", 2, sub.number)
    _render_legacy(document, sub)
    _render_blocks(document, sub.blocks)


def _recommendations(document: Document, spec: ReportSpecification) -> None:
    _heading(document, "11. RECOMENDACIONES", 1, "11")
    if not spec.recommendations:
        _notice(document, NO_EVIDENCE)
        return
    _paragraph(document, "Con base en las alarmas registradas durante el período, se destacan las siguientes "
                         "recomendaciones y conclusiones:")
    for finding in spec.recommendations:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _font(paragraph.add_run(finding.statement), size=11)
        if finding.evidence and finding.evidence.source_alarm_numbers:
            _font(paragraph.add_run(f" (Evidencia: {', '.join(finding.evidence.source_alarm_numbers)})"), size=9.5,
                  italic=True, color=MUTED)


def _update_fields_on_open(document: Document) -> None:
    element = OxmlElement("w:updateFields")
    element.set(qn("w:val"), "true")
    document.settings.element.append(element)


def build_document(spec: ReportSpecification, output_path: Path) -> Path:
    document = Document()
    _style_document(document)

    # Front matter: no header band; the cover has neither header nor footer.
    front = document.sections[0]
    front.different_first_page_header_footer = True
    _footer_box(front.footer, spec)
    _cover(document, spec)
    _table_of_contents(document, spec)
    _document_control(document, spec)
    _objective_and_services(document, spec)

    # Chapters: gray header band with the Invotecsa logo on every page.
    body = document.add_section(WD_SECTION.NEW_PAGE)
    body.different_first_page_header_footer = False
    body.top_margin = Inches(1.35)
    _header_band(body)
    body.footer.is_linked_to_previous = True

    for section in spec.sections:
        _render_section(document, section)
    _recommendations(document, spec)

    _update_fields_on_open(document)
    embed_font(document, HEADING_FONT, {"regular": "Poppins-Medium.ttf", "bold": "Poppins-Bold.ttf"})
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return output_path


__all__ = ["WD_BREAK", "build_document", "resolve_logo_path"]
