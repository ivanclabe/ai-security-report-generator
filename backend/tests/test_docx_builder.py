from docx import Document

from app.document.docx_builder import build_document
from app.schemas.report_spec import (
    DocumentControl,
    EvidenceRef,
    KeyFinding,
    ReportPeriod,
    ReportSpecification,
    Section,
    SignificantAlert,
    SimpleTable,
    Subsection,
)


def _sample_spec() -> ReportSpecification:
    return ReportSpecification(
        document_control=DocumentControl(titulo="Informe de Gestión — THOMAS TGS-TPS", fecha_version="2026-08-06"),
        company="THOMAS TGS-TPS",
        period=ReportPeriod(start="2026-07-01", end="2026-07-31"),
        objective="Dar a conocer el estado general de la gestión de seguridad.",
        services=["SIEM", "Antígeno de Red", "Antígeno de Correo"],
        sections=[
            Section(
                number="3",
                title="Inventario Monitoreado",
                summary="Resumen del inventario.",
                subsections=[
                    Subsection(
                        number="3.2",
                        title="Indicadores Alarmas",
                        table=SimpleTable(headers=["Riesgo", "Cantidad"], rows=[["Crítico", "19"], ["Alto", "59"]]),
                    ),
                ],
            ),
            Section(
                number="4",
                title="Gestión Alarmas SIEM",
                summary="Resumen SIEM.",
                significant_alerts=[
                    SignificantAlert(
                        alarm_number="ALARMA TGS TPS 3636",
                        heading="ALARMA TGS TPS 3636",
                        description="Se detectaron agentes desconectados.",
                        response="Se validó con el usuario.",
                    )
                ],
            ),
            Section(
                number="7",
                title="Inteligencia de Amenazas",
                subsections=[
                    Subsection(
                        number="7.2",
                        title="Vulnerabilidades",
                        available=False,
                        unavailable_reason="No disponible: requiere un feed de gestión de vulnerabilidades.",
                    )
                ],
            ),
            Section(
                number="8",
                title="Análisis de Comportamiento",
                available=False,
                unavailable_reason="No disponible: requiere una plataforma UEBA.",
            ),
        ],
        recommendations=[
            KeyFinding(
                statement="Reforzar la protección perimetral de la red.",
                evidence=EvidenceRef(source_alarm_numbers=["ALARMA TGS TPS 3636"], metric="15 alarmas"),
            )
        ],
    )


def test_build_document_creates_valid_docx_file(tmp_path):
    spec = _sample_spec()
    output_path = tmp_path / "informe.docx"

    result_path = build_document(spec, output_path)

    assert result_path.exists()
    assert result_path.stat().st_size > 0
    Document(str(result_path))  # raises if the file isn't a valid docx


def test_build_document_includes_section_headings(tmp_path):
    spec = _sample_spec()
    output_path = tmp_path / "informe.docx"
    build_document(spec, output_path)

    doc = Document(str(output_path))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]

    assert any("Objetivo" in h for h in headings)
    assert any("Servicios" in h for h in headings)
    assert any("3. Inventario Monitoreado" in h for h in headings)
    assert any("3.2 Indicadores Alarmas" in h for h in headings)
    assert any("Recomendaciones" in h for h in headings)


def test_build_document_renders_unavailable_sections_with_notice(tmp_path):
    spec = _sample_spec()
    output_path = tmp_path / "informe.docx"
    build_document(spec, output_path)

    doc = Document(str(output_path))
    full_text = "\n".join(p.text for p in doc.paragraphs)

    assert "requiere una plataforma UEBA" in full_text
    assert "requiere un feed de gestión de vulnerabilidades" in full_text


def test_build_document_renders_significant_alert_and_response(tmp_path):
    spec = _sample_spec()
    output_path = tmp_path / "informe.docx"
    build_document(spec, output_path)

    doc = Document(str(output_path))
    full_text = "\n".join(p.text for p in doc.paragraphs)

    assert "ALARMA TGS TPS 3636" in full_text
    assert "Se validó con el usuario." in full_text


def test_build_document_renders_table_rows(tmp_path):
    spec = _sample_spec()
    output_path = tmp_path / "informe.docx"
    build_document(spec, output_path)

    doc = Document(str(output_path))
    assert len(doc.tables) >= 1
    found = False
    for table in doc.tables:
        cell_texts = [c.text for row in table.rows for c in row.cells]
        if "Crítico" in cell_texts and "19" in cell_texts:
            found = True
    assert found


# ---- clean text: no highlighter anywhere ------------------------------------------


def _document_xml(path) -> str:
    import zipfile

    with zipfile.ZipFile(path) as z:
        return z.read("word/document.xml").decode("utf-8")


def test_generated_document_has_no_highlighted_text(tmp_path):
    """Regression: summary paragraphs, section/subsection numbers and the cover
    title used to carry a cyan/yellow highlight. The finished document must
    contain none, anywhere (body, tables, cover)."""
    out = build_document(_sample_spec(), tmp_path / "r.docx")
    xml = _document_xml(out)

    assert "<w:highlight" not in xml
    assert "w:shd" in xml  # table header shading (a cell fill, not text highlight) is intentional


def test_summary_paragraphs_are_plain_runs(tmp_path):
    out = build_document(_sample_spec(), tmp_path / "r.docx")
    doc = Document(str(out))
    summary = next(p for p in doc.paragraphs if p.text == "Resumen del inventario.")
    assert all(run.font.highlight_color is None for run in summary.runs)


# ---- per-company logo -------------------------------------------------------------


def test_logo_resolution_uses_the_companys_file_and_falls_back_to_the_default(tmp_path, monkeypatch):
    from app.document import docx_builder

    logos = tmp_path / "logos"
    logos.mkdir()
    (logos / "thomas-tps.png").write_bytes(b"png")
    monkeypatch.setattr(docx_builder, "LOGOS_DIR", logos)

    assert docx_builder.resolve_logo_path("Thomas TPS") == logos / "thomas-tps.png"
    assert docx_builder.resolve_logo_path("TGS BRASIL") == docx_builder.DEFAULT_LOGO_PATH  # file not supplied yet
    assert docx_builder.resolve_logo_path(None) == docx_builder.DEFAULT_LOGO_PATH


def test_logo_key_cannot_escape_the_logos_directory(tmp_path, monkeypatch):
    from app.document import docx_builder

    (tmp_path / "secret.png").write_bytes(b"png")
    logos = tmp_path / "logos"
    logos.mkdir()
    monkeypatch.setattr(docx_builder, "LOGOS_DIR", logos)

    assert docx_builder.resolve_logo_path("../secret") == docx_builder.DEFAULT_LOGO_PATH


def test_document_embeds_the_companys_logo_when_available(tmp_path, monkeypatch):
    from PIL import Image

    from app.document import docx_builder

    logos = tmp_path / "logos"
    logos.mkdir()
    Image.new("RGB", (40, 20), (200, 30, 30)).save(logos / "thomas-tps.png")
    monkeypatch.setattr(docx_builder, "LOGOS_DIR", logos)

    spec = _sample_spec().model_copy(update={"company_logo": "Thomas TPS"})
    out = build_document(spec, tmp_path / "tps.docx")

    import zipfile

    with zipfile.ZipFile(out) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
        sizes = {z.getinfo(n).file_size for n in media if n.endswith(".png")}
    assert (logos / "thomas-tps.png").stat().st_size in sizes
